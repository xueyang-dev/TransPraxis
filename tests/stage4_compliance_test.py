"""Stage 4 source-backed compliance and language constraints."""
from __future__ import annotations

import io

import core
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Cm, Pt
from transpraxis import compliance
from transpraxis import thesis_constraints


def _state(text="", **settings):
    return {"p3_md": text, "pairs": [], "research_settings": settings,
            "report_enabled": True}


def _report(abstract="摘" * 400, keywords=None, appendices=None):
    return {"report": {
        "abstract_zh": abstract,
        "keywords_zh": (keywords or {}).get("zh", ["a"] * 5),
        "keywords_en": (keywords or {}).get("en", ["b"] * 5),
        "appendices": appendices or [],
    }}


def test_default_profile_is_anonymous_and_source_backed():
    profile = compliance.compliance_profile()
    assert profile["profile_id"] == compliance.DEFAULT_PROFILE_ID
    assert profile["display_name"] == "默认 MTI 实践报告规范"
    assert profile["profile_type"] == "default_mti_practice_report"
    assert "institution" not in profile and "college" not in profile
    assert all(rule["enforcement"] in compliance.ENFORCEMENT
               for rule in profile["rules"])
    assert all(rule["source_document"] != "docs/mti-practice-driven-roadmap.md"
               for rule in profile["rules"])
    authority_rules = [rule for rule in profile["rules"]
                       if rule["authority_level"] != "project"]
    sourced_rules = [rule for rule in authority_rules
                     if rule["reliable_source_mapping"]]
    assert sourced_rules
    assert all(rule["source_id"] and rule["source_recorded"]
               for rule in sourced_rules)
    assert all(rule["authority_level"] == "reference_template"
               for rule in authority_rules)
    assert {source["source_id"] for source in profile["sources"]} == {
        compliance.REFERENCE_SOURCE_ID,
    }
    assert profile["implementation_sources"][0]["document"] == \
        "docs/mti-practice-driven-roadmap.md"
    assert all(rule["source_type"] == "reference_template"
               for rule in sourced_rules)
    assert all(not rule["source_url"] for rule in sourced_rules)
    assert all(rule["page_or_clause"] and rule["page_or_clause"] != "待提供"
               for rule in sourced_rules)


def test_rules_without_reliable_sources_cannot_be_enforced():
    profile = compliance.compliance_profile()
    profile["rules"].append({
        "rule_id": "unmapped_custom_rule", "category": "unsupported",
        "description": "Paper usually...", "authority_level": "custom_profile",
        "source_document": "docs/sources/not-present.md", "source_date": "2026",
        "page_or_clause": "page 1", "source_excerpt_or_summary": "claim",
        "source_available": False, "scope": "report", "check_type": "manual",
        "severity": "error", "enforcement": "enforced", "expected": "",
        "conflicts_with": [], "supersedes": [],
    })
    fake = profile["rules"][-1]
    assert fake["source_available"] is False
    normalized = compliance._result(fake, "manual_review", "source mapping unavailable")
    assert normalized["enforcement"] != "enforced"
    result = compliance.evaluate_compliance({}, {}, profile, "")
    assert not any(item["rule_id"] == "unmapped_custom_rule" and
                   item["enforcement"] == "enforced"
                   for item in result["rules"])


def test_reliable_custom_profile_mapping_can_remain_enforced(tmp_path, monkeypatch):
    source = tmp_path / "reference-template.md"
    source.write_text("clause", encoding="utf-8")
    monkeypatch.setattr(compliance, "SOURCE_ROOT", tmp_path)
    rule = compliance._rule(
        "sourced", "layout", "Sourced rule", authority_level="custom_profile",
        source_document=source.name, page_or_clause="clause 1",
        enforcement="enforced")
    assert rule["reliable_source_mapping"] is True
    assert rule["enforcement"] == "enforced"


def test_stage45_formal_rules_use_structured_sources_without_bundling_originals():
    profile = compliance.compliance_profile()
    expected_enforced = {
        "abstract_zh_length", "keywords_count", "toc_depth",
        "citation_reference_bidirectional", "figure_table_numbering",
        "bilingual_appendix", "case_conclusion_structure", "docx_layout",
    }
    rules = {rule["rule_id"]: rule for rule in profile["rules"]}
    assert expected_enforced <= {
        rule_id for rule_id, rule in rules.items()
        if rule["enforcement"] == "enforced"
    }
    assert all(rules[rule_id]["source_id"] and
               rules[rule_id]["reliable_source_mapping"] and
               rules[rule_id]["source_file_present"] is False
               for rule_id in expected_enforced)
    assert profile["authority_mapping_status"] == "reference_template_mapped"
    result = compliance.evaluate_compliance(_state(), _report(), profile)
    assert result["source_audit"]["enforced_rule_count"] == len(expected_enforced)
    assert result["source_audit"]["rules_without_source_mapping"] == []


def test_stage45_uncertain_rules_and_citation_conflict_remain_manual():
    profile = compliance.compliance_profile()
    rules = {rule["rule_id"]: rule for rule in profile["rules"]}
    assert rules["source_length"]["enforcement"] == "manual_review"
    assert rules["source_length"]["source_id"] == compliance.REFERENCE_SOURCE_ID
    assert rules["synthetic_case_policy"]["enforcement"] == "manual_review"
    assert rules["synthetic_case_policy"]["reliable_source_mapping"] is True
    conflict = profile["conflicts"][0]
    assert conflict["rule"] == "citation_style"
    assert {"footnote", "numeric_sequence"} <= set(conflict["options"])
    assert conflict["resolved_by"] == "manual_review"


def test_stage45_layout_values_are_checked_when_ooxml_facts_are_available():
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.3)
    section.bottom_margin = Cm(3.3)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.8)
    section.header_distance = Cm(2.6)
    section.footer_distance = Cm(2.6)
    paragraph = document.add_paragraph("正文")
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(20)
    stream = io.BytesIO()
    document.save(stream)

    result = compliance.evaluate_compliance(
        _state(), _report(), compliance.compliance_profile(),
        docx_bytes=stream.getvalue())
    rule = next(x for x in result["rules"] if x["rule_id"] == "docx_layout")
    assert rule["status"] == "pass"
    assert rule["enforcement"] == "enforced"


def test_chinese_abstract_boundary_and_statistics_are_explicit():
    profile = compliance.compliance_profile()
    for count, expected in ((399, "fail"), (400, "pass"), (600, "pass"),
                            (601, "fail")):
        result = compliance.evaluate_compliance(
            _state(), _report(abstract="摘" * count), profile)
        rule = next(x for x in result["rules"]
                    if x["rule_id"] == "abstract_zh_length")
        assert rule["status"] == expected
        assert rule["actual"] == count


def test_keywords_require_five_to_eight_and_both_languages():
    profile = compliance.compliance_profile()
    for zh, en, expected in ((4, 5, "fail"), (5, 8, "pass"), (9, 5, "fail")):
        result = compliance.evaluate_compliance(
            _state(), _report(keywords={"zh": ["x"] * zh, "en": ["y"] * en}),
            profile)
        assert next(x for x in result["rules"]
                    if x["rule_id"] == "keywords_count")["status"] == expected


def test_toc_depth_counts_visible_levels_not_prose_mentions():
    profile = compliance.compliance_profile()
    result = compliance.evaluate_compliance(
        _state("#### 四级正文"), _report(), profile)
    assert next(x for x in result["rules"]
                if x["rule_id"] == "toc_depth")["status"] == "fail"


def test_toc_missing_required_components_is_visible_when_toc_facts_exist():
    profile = compliance.compliance_profile()
    result = compliance.evaluate_compliance(
        _state("# 正文"), {
            "report": {"report": {"abstract_zh": "摘" * 400,
                                   "toc": {"expected": ["目录", "参考文献"],
                                            "actual": ["目录"]}}}},
        profile)
    rule = next(x for x in result["rules"] if x["rule_id"] == "toc_depth")
    assert rule["status"] == "fail"
    assert rule["actual"]["missing"] == ["参考文献"]


def test_citation_reference_bidirectional_and_duplicates():
    profile = compliance.compliance_profile()
    artifacts = {"literature_sources": {"sources": [
        {"source_id": "a"}, {"source_id": "b"}, {"source_id": "b"}]}}
    result = compliance.evaluate_compliance(
        _state("<!--cite:a-->"), artifacts, profile, "<!--cite:a-->")
    rule = next(x for x in result["rules"]
                if x["rule_id"] == "citation_reference_bidirectional")
    assert rule["status"] == "fail"
    assert rule["actual"]["duplicate_ids"] == ["b"]


def test_figure_and_table_numbering_checks_captions():
    profile = compliance.compliance_profile()
    text = "图 3.1：系统结构\n表 4.1 数据"
    result = compliance.evaluate_compliance(_state(text), _report(), profile, text)
    rule = next(x for x in result["rules"]
                if x["rule_id"] == "figure_table_numbering")
    assert rule["status"] == "pass"


def test_bilingual_appendix_roles():
    profile = compliance.compliance_profile()
    result = compliance.evaluate_compliance(
        _state(), _report(appendices=["附录一：原文与译文"]), profile)
    assert next(x for x in result["rules"]
                if x["rule_id"] == "bilingual_appendix")["status"] == "pass"


def test_formal_report_without_appendix_is_not_silently_not_applicable():
    profile = compliance.compliance_profile()
    result = compliance.evaluate_compliance(_state(), _report(), profile)
    assert next(x for x in result["rules"]
                if x["rule_id"] == "bilingual_appendix")["status"] == "fail"


def test_chinese_source_length_uses_layout_cjk_count():
    profile = compliance.compliance_profile()
    state = _state()
    state["paras"] = ["汉" * 10001]
    result = compliance.evaluate_compliance(state, _report(), profile, "")
    assert next(x for x in result["rules"]
                if x["rule_id"] == "source_length")["status"] == "pass"


def test_english_source_length_is_manual_until_conversion_rule_confirmed():
    profile = compliance.compliance_profile()
    state = _state()
    state["paras"] = ["word " * 11000]
    result = compliance.evaluate_compliance(state, _report(), profile, "")
    rule = next(x for x in result["rules"] if x["rule_id"] == "source_length")
    assert rule["status"] == "manual_review"
    assert "需要根据所在院校要求确认" in rule["message"]


def test_citation_conflict_is_manual_and_does_not_double_enforce():
    profile = compliance.compliance_profile()
    result = compliance.evaluate_compliance(_state(), _report(), profile, "")
    rules = {x["rule_id"]: x for x in result["rules"]}
    assert rules["citation_style"]["status"] == "manual_review"
    assert rules["citation_style"]["actual"] == "configurable"
    assert result["conflicts"][0]["resolved_by"] == "manual_review"


def test_project_constraints_stay_separate_from_default_profile():
    profile = compliance.compliance_profile()
    result = compliance.evaluate_compliance(_state(), _report(), profile, "")
    assert result["project_constraints"]["status"] in {"pass", "fail"}
    assert result["profile_compliance"]["enforced_rule_count"] == 8


def test_forbidden_phrase_location_is_exact():
    text = "# 3.3.2 策略\n\n这一段使用谱系对位。"
    result = compliance.evaluate_language_constraints(
        _state(text, forbidden_report_phrases=["谱系对位"]), text)
    assert result["status"] == "fail"
    assert result["failures"][0]["occurrences"][0]["section"] == \
        "3.3.2 策略"


def test_allowed_theory_labels_detect_unlisted_theory():
    text = "本文使用关联理论。"
    result = compliance.evaluate_language_constraints(
        _state(text, allowed_theory_labels=["目的论"]), text)
    assert result["status"] == "manual_review"
    assert "关联理论" in result["constraints"][-1]["value"]


def test_required_terminology_string_is_checked_and_reported():
    result = compliance.evaluate_language_constraints(
        _state("本文没有该词。", required_terminology=["目标术语"]),
        "本文没有该词。")
    missing = next(x for x in result["constraints"]
                   if x["kind"] == "required_terminology")
    assert missing["status"] == "manual_review"
    assert missing["value"] == ["目标术语"]


def test_placeholders_become_manual_review_items():
    profile = compliance.compliance_profile()
    text = "致谢：导师是【待作者填写】。"
    result = compliance.evaluate_compliance(_state(text), _report(), profile, text)
    rule = next(x for x in result["rules"]
                if x["rule_id"] == "author_placeholders")
    assert rule["status"] == "manual_review"
    assert rule["actual"][0]["excerpt"] == "致谢：导师是【待作者填写】。"


def test_synthetic_policy_separates_project_and_profile():
    profile = compliance.compliance_profile()
    artifacts = {"selected_cases": {
        "synthetic_count_policy": "counts_toward_minimum",
        "cases": [{"case_id": "SC-1", "case_type": "synthetic_contrast"}]}}
    result = compliance.evaluate_compliance(_state(), artifacts, profile, "")
    rule = next(x for x in result["rules"]
                if x["rule_id"] == "synthetic_case_policy")
    assert rule["status"] == "manual_review"
    assert rule["actual"]["project_constraint"] == "counts_toward_minimum"


def test_unknown_profile_identifier_normalizes_to_default(tmp_path):
    old_output = core.OUTPUT_DIR
    core.OUTPUT_DIR = tmp_path
    try:
        state = core.new_job_state("legacy.docx")
        state["compliance_profile_id"] = "PRIVATE_CUSTOM_PROFILE"
        result = core.compliance_profile_view("legacyjob", state)
        assert result["profile_id"] == compliance.DEFAULT_PROFILE_ID
        assert "language_constraints" in result
    finally:
        core.OUTPUT_DIR = old_output


def test_default_report_structure_is_generic_and_customizable():
    constraints = thesis_constraints.build_constraints({})
    assert constraints["structure_source"] == "default_mti_profile"
    assert [chapter["role"] for chapter in constraints["chapters"]] == [
        "introduction", "project_overview", "case_analysis",
        "conclusion_reflection"]
    assert [item["title"] for item in constraints["front_matter"][:2]] == [
        "中文摘要", "ABSTRACT"]
    assert [item["role"] for item in constraints["back_matter"]] == [
        "references", "appendix"]

    customized = thesis_constraints.build_constraints({
        "report_sections": [{"section_id": "A", "title": "自定义章节"}],
    })
    assert customized["structure_source"] == "user_configured_sections"
    assert [chapter["title"] for chapter in customized["chapters"]] == ["自定义章节"]
